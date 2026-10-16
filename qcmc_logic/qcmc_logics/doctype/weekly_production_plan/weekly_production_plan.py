from __future__ import annotations

from io import BytesIO

import frappe
from frappe import _
from frappe.desk.search import validate_and_sanitize_search_inputs
from frappe.model.document import Document
from frappe.utils import add_days, flt, getdate


PLAN_LENGTH_DAYS = 14


class WeeklyProductionPlan(Document):
	def autoname(self):
		self._set_computed_values()
		self.name = make_record_no(self.section, self.start_date)

	def validate(self):
		self._set_computed_values()
		validate_section_company(self.company, self.section)
		self._validate_scheduled_items()
		self._populate_scheduled_item_details()
		self._populate_machine_schedule()

	def _validate_scheduled_items(self):
		if is_extrusion_section(self.section):
			if self.get("item_1_schedule") or self.get("item_2_schedule"):
				frappe.throw(_("EXTRUSION plans can only use the Extrusion Item List."))
			for row in self.get("extrusion_item_schedule"):
				if row.roll_code and not item_has_active_bom_for_machine(row.roll_code, row.machine):
					frappe.throw(
						_("Extrusion Item List, row {0}: {1} must have an active submitted BOM for machine {2}.").format(
							row.idx, row.roll_code, frappe.bold(row.machine)
						)
					)
			return

		if self.get("extrusion_item_schedule"):
			frappe.throw(_("The Extrusion Item List is only available for the EXTRUSION section."))

		for row in self.get("machine_schedule"):
			for fieldname, label in (("second_item", _("FG 1")), ("third_item", _("FG 2"))):
				item_code = row.get(fieldname)
				if item_code and not item_has_active_bom_for_machine(item_code, row.machine):
					frappe.throw(
						_("Row {0}: {1} must be an FG item with an active submitted BOM for machine {2}.").format(
							row.idx, label, frappe.bold(row.machine)
						)
					)

		for table_field, label in (("item_1_schedule", _("FG 1")), ("item_2_schedule", _("FG 2"))):
			for row in self.get(table_field):
				if row.product_code and not item_has_active_bom_for_machine(row.product_code, row.machine):
					frappe.throw(
						_("{0}, row {1}: {2} must have an active submitted BOM for machine {3}.").format(
							label, row.idx, row.product_code, frappe.bold(row.machine)
						)
					)

	def _populate_scheduled_item_details(self):
		if is_extrusion_section(self.section):
			for row in self.get("extrusion_item_schedule"):
				details = get_item_production_details(row.roll_code, row.machine, check_permission=False)
				row.std_weight_kg = details.get("std_weight_gm")
				row.roll_description = details.get("product_description")
				row.trim_percent = details.get("trim_percent")
				row.soph = details.get("soph")
				values = calculate_extrusion_values(
					row.paid_time,
					row.vos,
					row.change_over,
					row.stabilization,
					row.soph,
					row.std_weight_kg,
				)
				row.update(values)
			return

		for table_field in ("item_1_schedule", "item_2_schedule"):
			for row in self.get(table_field):
				details = get_item_production_details(row.product_code, row.machine, check_permission=False)
				for fieldname in ("product_description", "cavs", "soph", "std_weight_gm", "cycle_time"):
					row.set(fieldname, details.get(fieldname))
				if table_field == "item_1_schedule":
					row.plan_hrs = flt(row.paid_hrs) - flt(row.vos) - flt(row.cm_co) - flt(row.stab)
					row.daily_volume_qty = row.plan_hrs * flt(row.soph)
				else:
					row.daily_volume_qty = 24 * flt(row.soph)
				row.daily_volume_kgs = flt(row.std_weight_gm) * flt(row.daily_volume_qty) / 1000

	def _set_computed_values(self):
		if not self.start_date:
			return

		self.end_date = add_days(self.start_date, PLAN_LENGTH_DAYS - 1)
		if self.section:
			self.record_no = make_record_no(self.section, self.start_date)

	def _populate_machine_schedule(self):
		if not self.start_date or not self.section:
			self.set("machine_schedule", [])
			return

		workstations = get_section_workstations(self.section)
		if not workstations:
			frappe.throw(
				_("No enabled Workstations are assigned to Plant Floor {0}.").format(
					frappe.bold(self.section)
				)
			)

		# Keep the planner's saved choices when validation rebuilds the generated
		# machine/date matrix.  The live Sales Order candidates are deliberately
		# not fields in the document and are fetched by the client when a row opens.
		saved_slots = {}
		for row in self.get("machine_schedule"):
			if not row.machine or not row.plan_date:
				continue
			saved_slots.setdefault((row.machine, getdate(row.plan_date)), []).append({
				"second_item": row.second_item,
				"second_item_description": row.second_item_description,
				"third_item": row.third_item,
				"third_item_description": row.third_item_description,
			})

		self.set("machine_schedule", [])
		for workstation in workstations:
			for day_offset in range(PLAN_LENGTH_DAYS):
				plan_date = add_days(self.start_date, day_offset)
				saved_rows = saved_slots.get((workstation, getdate(plan_date))) or [{}]
				for saved_row in saved_rows:
					self.append("machine_schedule", {
						"machine": workstation,
						"plan_date": plan_date,
						**saved_row,
					})


def make_record_no(section: str, start_date: str) -> str:
	"""Build IDs like CPS101926 from a section and 2026-10-19."""
	if not section or not start_date:
		frappe.throw(_("Section and Start Date are required to create the Record No."))

	section_code = "".join(character for character in section.upper() if character.isalnum())
	return f"{section_code}{getdate(start_date).strftime('%m%d%y')}"


def is_extrusion_section(section: str | None) -> bool:
	return (section or "").strip().upper().startswith("EXTRUSION")


def calculate_extrusion_values(
	paid_time: float,
	vos: float,
	change_over: float,
	stabilization: float,
	soph: float,
	std_weight: float,
) -> dict[str, float]:
	"""Calculate output for a roll/extrusion production-plan row."""
	plan_hours = flt(paid_time) - flt(vos) - flt(change_over) - flt(stabilization)
	number_of_kgs = plan_hours * flt(soph)
	weight = flt(std_weight)

	return {
		"plan_hours": plan_hours,
		"number_of_kgs": number_of_kgs,
		"number_of_rolls": number_of_kgs / weight if weight else 0,
	}


def get_section_workstations(section: str) -> list[str]:
	return frappe.get_all(
		"Workstation",
		filters={"plant_floor": section, "disabled": 0},
		pluck="name",
		order_by="name asc",
	)


def _normalize_excel_header(value) -> str:
	return "".join(character for character in str(value or "").strip().lower() if character.isalnum())


def extract_production_rows(worksheet) -> list[dict]:
	"""Read Product Code and For Production from a worksheet with flexible header placement."""
	product_aliases = {"productcode", "prodcode", "itemcode"}
	quantity_aliases = {"forproduction", "productionqty"}
	header_row = product_column = quantity_column = None

	for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
		headers = [_normalize_excel_header(value) for value in values]
		product_matches = [index for index, header in enumerate(headers) if header in product_aliases]
		quantity_matches = [
			index for index, header in enumerate(headers)
			if header in quantity_aliases or header.startswith("forproduction")
		]
		if product_matches and quantity_matches:
			header_row, product_column, quantity_column = row_number, product_matches[0], quantity_matches[0]
			break

	if header_row is None:
		frappe.throw(_("Excel must contain Product Code and For Production columns."))

	quantities = {}
	for excel_row, values in enumerate(
		worksheet.iter_rows(min_row=header_row + 1, values_only=True),
		start=header_row + 1,
	):
		product_code = str(values[product_column] or "").strip() if product_column < len(values) else ""
		raw_quantity = values[quantity_column] if quantity_column < len(values) else None
		if not product_code or raw_quantity in (None, "", "-"):
			continue
		try:
			quantity = flt(float(str(raw_quantity).replace(",", "").strip()))
		except (TypeError, ValueError):
			frappe.throw(_("Invalid For Production quantity for product {0}: {1}").format(product_code, raw_quantity))
		if quantity > 0:
			entry = quantities.setdefault(product_code, {"quantity": 0, "excel_rows": []})
			entry["quantity"] += quantity
			entry["excel_rows"].append(excel_row)

	return [
		{
			"product_code": code,
			"for_production_qty": values["quantity"],
			"excel_rows": values["excel_rows"],
		}
		for code, values in quantities.items()
	]


@frappe.whitelist()
def import_production_plan(
	file_url: str,
	section: str,
	start_date: str,
	occupied_slots=None,
	plan_name: str | None = None,
) -> dict:
	"""Parse an uploaded workbook and allocate its production quantities to free machine/day slots."""
	if not frappe.has_permission("Weekly Production Plan", "write"):
		frappe.throw(_("Not permitted to update Weekly Production Plans."), frappe.PermissionError)
	if plan_name:
		plan = frappe.get_doc("Weekly Production Plan", plan_name)
		plan.check_permission("write")
		if plan.docstatus != 0:
			frappe.throw(_("Production Plan Excel can only be imported while the plan is Active."))
	if not file_url or not section or not start_date:
		frappe.throw(_("Production Plan Excel, Section, and Start Date are required."))
	if is_extrusion_section(section):
		frappe.throw(_("This Product Code / For Production importer is only available for FG plans."))
	if not file_url.lower().endswith((".xlsx", ".xlsm")):
		frappe.throw(_("Please upload an .xlsx or .xlsm Excel workbook."))

	from openpyxl import load_workbook

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	file_doc.check_permission("read")
	workbook = load_workbook(BytesIO(file_doc.get_content()), data_only=True, read_only=True)
	production_rows = extract_production_rows(workbook.active)
	workstations = get_section_workstations(section)
	occupied = {
		(str(row.get("machine")), str(getdate(row.get("plan_date"))))
		for row in frappe.parse_json(occupied_slots or "[]")
		if row.get("machine") and row.get("plan_date")
	}

	assignments, unassigned = [], []
	scheduled_products = 0
	for production_row in production_rows:
		item_code = production_row["product_code"]
		eligible = [machine for machine in workstations if item_has_active_bom_for_machine(item_code, machine)]
		eligible.sort(
			key=lambda machine: not item_has_active_bom_for_machine(item_code, machine, default_only=True)
		)
		remaining_quantity = flt(production_row["for_production_qty"])
		next_day_offset = 0
		product_was_scheduled = False
		details_by_machine = {}

		while remaining_quantity > 0:
			selected = None
			for day_offset in range(next_day_offset, PLAN_LENGTH_DAYS):
				plan_date = add_days(start_date, day_offset)
				for machine in eligible:
					if (machine, str(getdate(plan_date))) in occupied:
						continue
					if machine not in details_by_machine:
						details_by_machine[machine] = get_item_production_details(
							item_code, machine, check_permission=False
						)
					details = details_by_machine[machine]
					if flt(details.get("soph")) > 0:
						selected = (machine, plan_date, day_offset, details)
						break
				if selected:
					break

			if not selected:
				unassigned.append({
					"product_code": item_code,
					"for_production_qty": remaining_quantity,
					"excel_rows": production_row["excel_rows"],
					"reason": _("No compatible free machine/day with a positive SOPH"),
				})
				break

			machine, plan_date, day_offset, details = selected
			soph = flt(details.get("soph"))
			daily_quantity = min(remaining_quantity, 24 * soph)
			paid_hours = daily_quantity / soph
			occupied.add((machine, str(getdate(plan_date))))
			assignments.append({
				"product_code": item_code,
				"imported_from_excel": 1,
				"machine": machine,
				"plan_date": plan_date,
				"paid_hrs": paid_hours,
				"plan_hrs": paid_hours,
				"daily_volume_qty": daily_quantity,
				"daily_volume_kgs": flt(details.get("std_weight_gm")) * daily_quantity / 1000,
				"f_rate": 100,
				**{key: details.get(key) for key in ("product_description", "cavs", "soph", "std_weight_gm", "cycle_time")},
			})
			product_was_scheduled = True
			remaining_quantity = max(0, remaining_quantity - daily_quantity)
			next_day_offset = day_offset + 1

		if product_was_scheduled:
			scheduled_products += 1

	return {
		"rows": assignments,
		"unassigned": unassigned,
		"source_rows": len(production_rows),
		"scheduled_products": scheduled_products,
	}


def validate_section_company(company: str, section: str) -> None:
	if not company or not section:
		return

	section_company = frappe.db.get_value("Plant Floor", section, "company")
	if not section_company:
		frappe.throw(
			_("Plant Floor {0} does not have a Company assigned.").format(
				frappe.bold(section)
			)
		)
	if section_company != company:
		frappe.throw(
			_("Plant Floor {0} belongs to Company {1}, not {2}.").format(
				frappe.bold(section), frappe.bold(section_company), frappe.bold(company)
			)
		)


def item_has_active_bom_for_machine(item_code: str, machine: str, default_only: bool = False) -> bool:
	"""Match either the BOM header machine or a workstation on its operation rows."""
	return bool(
		frappe.db.sql(
			"""
			select bom.name
			from `tabBOM` bom
			where bom.item = %(item_code)s
				and bom.docstatus = 1
				and bom.is_active = 1
				and (%(default_only)s = 0 or bom.is_default = 1)
				and (
					bom.custom_machine = %(machine)s
					or exists (
						select 1
						from `tabBOM Operation` operation
						where operation.parent = bom.name
							and operation.parenttype = 'BOM'
							and operation.workstation = %(machine)s
					)
				)
			limit 1
			""",
			{"item_code": item_code, "machine": machine, "default_only": int(default_only)},
		)
	)


@frappe.whitelist()
def get_item_production_details(item_code: str, machine: str, check_permission: bool = True) -> dict:
	if check_permission and not frappe.has_permission("Weekly Production Plan", "read"):
		frappe.throw(_("Not permitted to read Weekly Production Plans."), frappe.PermissionError)
	if not item_code or not machine:
		return {}

	bom = frappe.db.sql(
		"""
		select
			bom.name,
			bom.custom_number_of_cavity as cavs,
			bom.custom_soph as soph,
			bom.custom_rate_per_minute as cycle_time,
			bom.custom_roll_trimmings_ as trim_percent
		from `tabBOM` bom
		where bom.item = %(item_code)s
			and bom.docstatus = 1
			and bom.is_active = 1
			and (
				bom.custom_machine = %(machine)s
				or exists (
					select 1 from `tabBOM Operation` operation
					where operation.parent = bom.name
						and operation.parenttype = 'BOM'
						and operation.workstation = %(machine)s
				)
			)
		order by bom.is_default desc, bom.modified desc
		limit 1
		""",
		{"item_code": item_code, "machine": machine},
		as_dict=True,
	)
	if not bom:
		return {}

	item = frappe.db.get_value(
		"Item", item_code, ["item_name", "description", "weight_per_unit"], as_dict=True
	) or {}
	return {
		"bom": bom[0].name,
		"product_description": item.get("description") or item.get("item_name"),
		"cavs": bom[0].cavs,
		"soph": bom[0].soph,
		"std_weight_gm": item.get("weight_per_unit"),
		"cycle_time": bom[0].cycle_time,
		"trim_percent": bom[0].trim_percent,
	}


@frappe.whitelist()
def get_machine_schedule(company: str, section: str, start_date: str) -> dict:
	if not company or not section or not start_date:
		return {"end_date": None, "record_no": None, "rows": []}

	validate_section_company(company, section)
	workstations = get_section_workstations(section)
	rows = [
		{"machine": workstation, "plan_date": add_days(start_date, day_offset)}
		for workstation in workstations
		for day_offset in range(PLAN_LENGTH_DAYS)
	]
	return {
		"end_date": add_days(start_date, PLAN_LENGTH_DAYS - 1),
		"record_no": make_record_no(section, start_date),
		"rows": rows,
	}


@frappe.whitelist()
def get_possible_sales_order_items(machine: str) -> list[dict]:
	"""Return live, unsaved SO candidates whose FG BOM is assigned to a machine."""
	if not frappe.has_permission("Weekly Production Plan", "read"):
		frappe.throw(_("Not permitted to read Weekly Production Plans."), frappe.PermissionError)
	if not machine:
		return []

	return frappe.db.sql(
		"""
		select
			so.name as sales_order,
			so.transaction_date as sales_order_date,
			so.docstatus,
			so.status,
			soi.item_code,
			soi.qty - coalesce(soi.delivered_qty, 0) as remaining_qty,
			coalesce(nullif(item.description, ''), item.item_name) as description
		from `tabSales Order` so
		inner join `tabSales Order Item` soi on soi.parent = so.name
		inner join `tabItem` item on item.name = soi.item_code
		where so.docstatus < 2
			and (so.docstatus = 0 or so.status in ('To Deliver and Bill', 'To Deliver'))
			and item.disabled = 0
			and coalesce(soi.delivered_qty, 0) < soi.qty
			and exists (
				select 1
				from `tabBOM` bom
				where bom.item = soi.item_code
					and bom.docstatus = 1
					and bom.is_active = 1
					and (
						bom.custom_machine = %(machine)s
						or exists (
							select 1
							from `tabBOM Operation` operation
							where operation.parent = bom.name
								and operation.parenttype = 'BOM'
								and operation.workstation = %(machine)s
						)
					)
			)
		order by so.transaction_date asc, so.name asc, soi.idx asc
		""",
		{"machine": machine},
		as_dict=True,
	)


@frappe.whitelist()
@validate_and_sanitize_search_inputs
def get_machine_fg_items(
	doctype: str,
	txt: str,
	searchfield: str,
	start: int,
	page_len: int,
	filters: dict | None = None,
) -> list[list]:
	"""Link-field options limited to FG items with a BOM for the row's machine."""
	machine = (filters or {}).get("machine")
	if not machine:
		return []

	return frappe.db.sql(
		"""
		select distinct item.name, item.description
		from `tabItem` item
		inner join `tabBOM` bom on bom.item = item.name
		where bom.docstatus = 1
			and bom.is_active = 1
			and item.disabled = 0
			and (
				bom.custom_machine = %(machine)s
				or exists (
					select 1
					from `tabBOM Operation` operation
					where operation.parent = bom.name
						and operation.parenttype = 'BOM'
						and operation.workstation = %(machine)s
				)
			)
			and (
				item.name like %(txt)s
				or item.item_name like %(txt)s
				or item.description like %(txt)s
			)
		order by item.name
		limit %(start)s, %(page_len)s
		""",
		{
			"machine": machine,
			"txt": f"%{txt}%",
			"start": start,
			"page_len": page_len,
		},
	)
